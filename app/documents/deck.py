from pathlib import Path

import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches
from sqlalchemy.orm import Session

from app.domain.analytics import AnalyticsService
from app.domain.inventory import InventoryService
from app.domain.preferences import PreferenceService


def generate_analysis_deck(session: Session, owner_id: str, days: int, output_dir: Path) -> Path:
    summary = AnalyticsService(session).summary(days)
    prefs = PreferenceService(session).all(owner_id)
    shop_name = prefs.get("shop_name", "Kirana Store")
    low_stock = InventoryService(session).low_stock()

    output_dir.mkdir(parents=True, exist_ok=True)
    chart_dir = output_dir / "charts"
    chart_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / f"sales_analysis_{days}d.pptx"

    prs = Presentation()
    title = prs.slides.add_slide(prs.slide_layouts[0])
    title.shapes.title.text = f"{shop_name} — Sales Analysis"
    title.placeholders[1].text = f"Last {days} day(s) | Generated from finalized store transactions"

    overview = prs.slides.add_slide(prs.slide_layouts[1])
    overview.shapes.title.text = "Business overview"
    overview.placeholders[1].text = (
        f"Sales: ₹{summary['total_sales']}\n"
        f"Bills: {summary['bill_count']}\n"
        f"GST collected: ₹{summary['tax_collected']}\n"
        f"Low-stock SKUs: {len(low_stock)}"
    )

    modes = summary["payment_modes"]
    if modes:
        labels = list(modes)
        values = [float(modes[k]) for k in labels]
        payment_chart = chart_dir / "payment_modes.png"
        fig, ax = plt.subplots(figsize=(7, 4))
        ax.bar(labels, values)
        ax.set_title("Sales by payment mode")
        ax.set_ylabel("INR")
        fig.tight_layout()
        fig.savefig(payment_chart, dpi=160)
        plt.close(fig)

        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "Payment mix"
        slide.shapes.add_picture(str(payment_chart), Inches(1.1), Inches(1.7), width=Inches(8.0))

    top_items = summary["top_items"]
    if top_items:
        labels = [row["name"] for row in top_items[:8]][::-1]
        values = [float(row["sales"]) for row in top_items[:8]][::-1]
        top_chart = chart_dir / "top_items.png"
        fig, ax = plt.subplots(figsize=(8, 4.5))
        ax.barh(labels, values)
        ax.set_title("Top items by sales")
        ax.set_xlabel("INR")
        fig.tight_layout()
        fig.savefig(top_chart, dpi=160, bbox_inches="tight")
        plt.close(fig)

        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "Top-selling products"
        slide.shapes.add_picture(str(top_chart), Inches(0.9), Inches(1.5), width=Inches(8.5))

    stock_slide = prs.slides.add_slide(prs.slide_layouts[1])
    stock_slide.shapes.title.text = "Stock health"
    if low_stock:
        stock_slide.placeholders[1].text = "\n".join(
            f"{p.name}: {p.stock_quantity:g} {p.base_unit} remaining (reorder at {p.reorder_level:g})"
            for p in low_stock[:12]
        )
    else:
        stock_slide.placeholders[1].text = "No products are at or below their reorder level."

    insight_slide = prs.slides.add_slide(prs.slide_layouts[1])
    insight_slide.shapes.title.text = "Operational insights"
    insights = []
    if top_items:
        insights.append(f"Highest-sales item: {top_items[0]['name']} (₹{top_items[0]['sales']}).")
    if modes:
        leading_mode = max(modes, key=lambda key: float(modes[key]))
        insights.append(f"Largest payment channel: {leading_mode} (₹{modes[leading_mode]}).")
    insights.append(f"{len(low_stock)} SKU(s) currently need reorder attention.")
    insight_slide.placeholders[1].text = "\n".join(insights)

    prs.save(path)
    return path
